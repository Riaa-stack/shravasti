import os
import uuid
from datetime import datetime
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from pptx import Presentation

from app.core.config import settings

def get_export_dir():
    export_dir = getattr(settings, "EXPORTS_STORAGE_DIR", "/data/exports")
    try:
        os.makedirs(export_dir, exist_ok=True)
    except Exception:
        export_dir = os.path.abspath(os.path.join(".", "data", "exports"))
        os.makedirs(export_dir, exist_ok=True)
    return export_dir

def export_to_pdf(content: str, filename: str) -> str:
    filepath = os.path.join(get_export_dir(), f"{filename}.pdf")
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )
    styles = getSampleStyleSheet()
    story = []

    # Title style
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=18,
        leading=22,
        textColor="#1a73e8"
    )
    story.append(Paragraph("ScholarForge AI Research Report", title_style))
    story.append(Spacer(1, 10))

    meta_style = ParagraphStyle(
        "CustomMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor="#5f6368"
    )
    story.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} &bull; ScholarForge AI", meta_style))
    story.append(Spacer(1, 16))

    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["BodyText"],
        fontSize=10,
        leading=14,
        textColor="#202124"
    )
    h2_style = ParagraphStyle(
        "CustomH2",
        parent=styles["Heading2"],
        fontSize=13,
        leading=17,
        textColor="#1a73e8"
    )
    h3_style = ParagraphStyle(
        "CustomH3",
        parent=styles["Heading3"],
        fontSize=11,
        leading=15,
        textColor="#202124"
    )

    for line in content.split("\n"):
        clean = line.strip()
        if not clean:
            story.append(Spacer(1, 6))
            continue
        if clean.startswith("# "):
            story.append(Paragraph(clean[2:], title_style))
            story.append(Spacer(1, 8))
        elif clean.startswith("## "):
            story.append(Paragraph(clean[3:], h2_style))
            story.append(Spacer(1, 6))
        elif clean.startswith("### "):
            story.append(Paragraph(clean[4:], h3_style))
            story.append(Spacer(1, 4))
        else:
            safe = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body_style))
            story.append(Spacer(1, 3))

    doc.build(story)
    return filepath

def export_to_docx(content: str, filename: str) -> str:
    filepath = os.path.join(get_export_dir(), f"{filename}.docx")
    document = Document()
    document.add_heading('ScholarForge AI Report', 0)
    
    for line in content.split('\n'):
        if line.strip():
            document.add_paragraph(line)
            
    document.save(filepath)
    return filepath

def export_to_pptx(content: str, filename: str) -> str:
    filepath = os.path.join(get_export_dir(), f"{filename}.pptx")
    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "ScholarForge AI Report"
    subtitle.text = f"Generated on {datetime.now().strftime('%Y-%m-%d')}"
    
    bullet_slide_layout = prs.slide_layouts[1]
    
    # Group lines by ~5 per slide
    lines = [l for l in content.split('\n') if l.strip()]
    
    for i in range(0, len(lines), 5):
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = "Report Content"
        tf = body_shape.text_frame
        
        for j, line in enumerate(lines[i:i+5]):
            if j == 0:
                tf.text = line
            else:
                p = tf.add_paragraph()
                p.text = line
                
    prs.save(filepath)
    return filepath
